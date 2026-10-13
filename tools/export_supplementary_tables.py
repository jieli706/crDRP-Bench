"""Export supplied supplementary table values without editing the workbook."""

import argparse
import csv
import datetime as dt
import hashlib
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]


def csv_value(value):
    if value is None:
        return ""
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    return str(value)


def table_rows(sheet):
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    while rows and not any(v is not None for v in rows[-1]):
        rows.pop()
    if len(rows) < 2:
        raise ValueError(f"Missing description/header rows in {sheet.title}")
    description = next((v for v in rows[0] if v is not None), "")
    data = rows[1:]
    width = max((i + 1 for row in data for i, v in enumerate(row) if v is not None), default=0)
    if not width:
        raise ValueError(f"Empty table in {sheet.title}")
    return str(description), [[csv_value(v) for v in row[:width]] + [""] * max(0, width - len(row)) for row in data]


def export_tables(workbook_path, output_dir):
    workbook_path = Path(workbook_path)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=False)
    if any(cell.data_type == "f" for sheet in workbook for row in sheet.iter_rows() for cell in row):
        workbook.close()
        raise ValueError("Workbook contains formulas. This value-only exporter requires a reviewed value workbook.")
    index = []
    for sheet in workbook:
        if sheet.title == "README":
            continue
        if not re.fullmatch(r"[A-Za-z0-9_-]+", sheet.title):
            raise ValueError(f"Unsafe sheet filename: {sheet.title}")
        description, rows = table_rows(sheet)
        target = output_dir / (sheet.title + ".csv")
        with target.open("w", encoding="utf-8", newline="") as stream:
            csv.writer(stream, lineterminator="\n").writerows(rows)
        index.append({"sheet": sheet.title, "csv_file": target.name, "data_rows": len(rows) - 1,
                      "columns": len(rows[0]), "description": description,
                      "sha256": hashlib.sha256(target.read_bytes()).hexdigest()})
    workbook.close()
    with (output_dir / "table_index.csv").open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=["sheet", "csv_file", "data_rows", "columns", "description", "sha256"], lineterminator="\n")
        writer.writeheader()
        writer.writerows(index)
    return index


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=ROOT / "data/Supplementary_Tables_S1-S15_REV7.xlsx")
    parser.add_argument("--output-dir", type=Path, default=ROOT / "out/revision/tables")
    args = parser.parse_args()
    index = export_tables(args.workbook, args.output_dir)
    print(json.dumps({"exported_tables": len(index), "total_data_rows": sum(x["data_rows"] for x in index), "workbook_unchanged": True}, indent=2))


if __name__ == "__main__":
    main()
